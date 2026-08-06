"""
Excel Generator Module
Project: MediTracker

Generates Excel Workbooks (.xlsx) using OpenPyXL.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelGenerator:

    @staticmethod
    def _apply_header_style(ws, col_count, fill_color="212529"):
        """Utility method to apply modern styles to column headers."""
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, col_count + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = alignment

    @staticmethod
    def _auto_adjust_columns(ws):
        """Auto-adjust column widths based on cell text lengths."""
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def generate_inventory_excel(cls, medicines, output_path):
        """
        Creates an Excel report for all active medicines.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory Summary"

        headers = ["ID", "Medicine Name", "Batch Number", "Category", "Quantity", "Price (₹)", "Total Value (₹)", "Expiry Date", "Status"]
        ws.append(headers)

        cls._apply_header_style(ws, len(headers), fill_color="1F4E78")

        for med in medicines:
            total_val = (med.quantity or 0) * (med.price or 0.0)
            status = "Expired" if med.is_expired else ("Low Stock" if med.is_low_stock else "Active")

            ws.append([
                med.id,
                med.name,
                med.batch_number,
                med.category or "Uncategorized",
                med.quantity,
                med.price,
                total_val,
                str(med.expiry_date) if med.expiry_date else "N/A",
                status
            ])

        cls._auto_adjust_columns(ws)
        wb.save(output_path)
        return output_path

    @classmethod
    def generate_sales_excel(cls, sales, output_path):
        """
        Creates an Excel report for sales history.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales Report"

        headers = ["Invoice #", "Customer Name", "Customer Phone", "Medicine Name", "Quantity", "Selling Price (₹)", "Total Amount (₹)", "Sale Date"]
        ws.append(headers)

        cls._apply_header_style(ws, len(headers), fill_color="2E75B6")

        for sale in sales:
            ws.append([
                sale.invoice_number or f"INV-{sale.id}",
                sale.customer_name or "Walk-in Customer",
                sale.customer_phone or "N/A",
                sale.medicine.name if sale.medicine else "Deleted Item",
                sale.quantity,
                sale.sale_price,
                sale.total_amount,
                sale.sale_date.strftime("%Y-%m-%d %H:%M") if sale.sale_date else "N/A"
            ])

        cls._auto_adjust_columns(ws)
        wb.save(output_path)
        return output_path